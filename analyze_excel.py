import openpyxl
import os

file_path = r"c:/Users/ThomasKhannaKlelundJ/OneDrive - Lauridsen Handel-Import A S/Dokumenter/Anti gravity/Test/Asfalt beregner/Asfalt_Beregning.xlsx"

try:
    print(f"Loading {file_path}...")
    wb = openpyxl.load_workbook(file_path, data_only=False)
    # print(f"All Sheets: {wb.sheetnames}")
    sheet = wb["Ark1"]
    print(f"Reading Sheet: {sheet.title}")
    
    with open("ark1_dump.txt", "w", encoding="utf-8") as f:
        f.write(f"Reading Sheet: {sheet.title}\n")
        f.write("-" * 30 + "\n")
        f.write("CELLS WITH VALUES/FORMULAS:\n")
        f.write("-" * 30 + "\n")
        
        rows = list(sheet.iter_rows(values_only=False))
        for row in rows:
            for cell in row:
                if cell.value:
                    f.write(f"Cell {cell.coordinate}: {cell.value}\n")
    print("Done writing to ark1_dump.txt")
                
except Exception as e:
    print(f"Error: {e}")
